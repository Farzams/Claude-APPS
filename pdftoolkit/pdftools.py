"""Extra PDF tools (no GUI): organize, crop, blank pages, stamps, compare, export, security, redaction, editing."""
from __future__ import annotations

import base64
import difflib
import html as _html
import io
import os
import re
import tempfile
from datetime import date
from pathlib import Path
from typing import Callable

import pymupdf

import core

Progress = Callable[[int, int, str], None] | None


# ---------------------------------------------------------------- helpers

def open_pdf(src: str, password: str = "") -> pymupdf.Document:
    doc = pymupdf.open(src)
    if doc.needs_pass and not doc.authenticate(password):
        doc.close()
        raise ValueError("The PDF is password protected" if not password else "Wrong password")
    return doc


def _save(doc: pymupdf.Document, dst: str, **kw):
    kw.setdefault("garbage", 3)
    kw.setdefault("deflate", True)
    doc.save(dst, **kw)


def _pages(spec: str, n: int) -> list[int]:
    return core.parse_pages(spec, n) if spec.strip() else list(range(n))


def _rgb(color) -> tuple[float, float, float]:
    if isinstance(color, str):
        c = color.lstrip("#")
        return tuple(int(c[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return tuple(color)


def _arch() -> pymupdf.Archive:
    d = core.fonts_dir()
    return pymupdf.Archive(str(d)) if d.is_dir() else pymupdf.Archive()


_FONT_CSS = ("@font-face{font-family:ArabicUI;src:url(NotoSansArabic-Regular.ttf)}"
             "@font-face{font-family:ArabicUI;font-weight:bold;src:url(NotoSansArabic-Bold.ttf)}"
             "@font-face{font-family:LatinUI;src:url(DejaVuSans.ttf)}"
             "@font-face{font-family:LatinUI;font-weight:bold;src:url(DejaVuSans-Bold.ttf)}")


def _text_css(size: float, color, bold=False, align="left", rtl=False) -> str:
    r, g, b = (int(v * 255) for v in _rgb(color))
    return (_FONT_CSS + f"body{{font-family:ArabicUI,LatinUI,sans-serif;font-size:{size}pt;color:rgb({r},{g},{b});"
            f"text-align:{align};margin:0;padding:0;{'font-weight:bold;' if bold else ''}"
            f"{'direction:rtl;' if rtl else ''}}}")


def put_text(page: pymupdf.Page, rect: pymupdf.Rect, text: str, size=11, color=(0, 0, 0), bold=False,
             align="left", opacity=1.0, overlay=True) -> bool:
    """Write (shaped, RTL-aware) text into a rectangle. Returns True when it fitted."""
    rtl = core.is_rtl_text(text)
    esc = _html.escape(text).replace("\n", "<br>")
    # MuPDF's HTML engine does not fall back between fonts per glyph, so pick the font per run.
    run = re.compile(r"((?:[\u0600-\u06FF\u0750-\u077F\uFB50-\uFDFF\uFE70-\uFEFF\u200c\u200f]+[ \u200c]*)+)")
    esc = "".join(f'<span style="font-family:ArabicUI">{p}</span>' if k % 2 else
                  (f'<span style="font-family:LatinUI">{p}</span>' if p else "")
                  for k, p in enumerate(run.split(esc)))
    css = _text_css(size, color, bold, "right" if rtl and align == "left" else align, rtl)
    spare, _ = page.insert_htmlbox(rect, esc, css=css, archive=_arch(), opacity=opacity, overlay=overlay,
                                   scale_low=0.4)
    return spare >= 0


def _tokens(text: str, i: int, n: int, name: str, bates: tuple[str, int, int] | None = None) -> str:
    out = (text.replace("{page}", str(i + 1)).replace("{pages}", str(n)).replace("{file}", name)
           .replace("{date}", date.today().isoformat()))
    if bates:
        prefix, start, digits = bates
        out = out.replace("{bates}", f"{prefix}{start + i:0{digits}d}")
    return out


def find_rects(page: pymupdf.Page, term: str) -> list[pymupdf.Rect]:
    """Rectangles of every match of `term`; works for Persian/Arabic (ignores spaces, half-spaces, letter forms)."""
    for v in core._search_variants(term):
        hits = page.search_for(v)
        if hits:
            return list(hits)
    target = core._fold(term)
    if not target:
        return []
    chars: list[tuple[str, pymupdf.Rect]] = []
    for b in page.get_text("rawdict")["blocks"]:
        for ln in b.get("lines", []):
            for sp in ln["spans"]:
                for ch in sp["chars"]:
                    f = core._fold(ch["c"])
                    for c in f:
                        chars.append((c, pymupdf.Rect(ch["bbox"])))
    text = "".join(c for c, _ in chars)
    rects: list[pymupdf.Rect] = []
    for tgt, seq in ((target, chars), (target[::-1], chars)):
        start = 0
        while True:
            at = text.find(tgt, start)
            if at < 0:
                break
            line = None
            for _, r in seq[at:at + len(tgt)]:
                if line is not None and abs(r.y0 - line.y0) < r.height * 0.6:
                    line |= r
                else:
                    if line is not None:
                        rects.append(line)
                    line = pymupdf.Rect(r)
            if line is not None:
                rects.append(line)
            start = at + len(tgt)
        if rects:
            break
    return rects


# ---------------------------------------------------------------- organize pages

def organize(src: str, dst: str, plan: list[tuple[int, int]]) -> int:
    """plan = [(source_page_index, extra_rotation_degrees), ...] in the wanted final order.

    Pages may repeat (duplicates); pages left out are deleted.
    """
    if not plan:
        raise ValueError("The result would have no pages")
    with open_pdf(src) as doc:
        rot = {}
        order = []
        for k, (p, r) in enumerate(plan):
            order.append(p)
            rot[k] = r
        doc.select(order)
        for k, r in rot.items():
            if r % 360:
                pg = doc[k]
                pg.set_rotation((pg.rotation + r) % 360)
        _save(doc, dst)
        return doc.page_count


def insert_pages(src: str, dst: str, other: str, at: int, pages: str = "") -> int:
    """Insert pages of another PDF (or image) before page index `at` (0-based; page_count = at the end)."""
    with open_pdf(src) as doc:
        if Path(other).suffix.lower() == ".pdf":
            with open_pdf(other) as o:
                sel = _pages(pages, o.page_count)
                for k, p in enumerate(sel):
                    doc.insert_pdf(o, from_page=p, to_page=p, start_at=at + k)
        else:
            tmp = os.path.join(tempfile.gettempdir(), f"pdftk_ins_{os.getpid()}.pdf")
            core.images_to_pdf([other], tmp)
            with pymupdf.open(tmp) as o:
                doc.insert_pdf(o, start_at=at)
            os.remove(tmp)
        _save(doc, dst)
        return doc.page_count


def extract_pages(src: str, dst: str, spec: str) -> int:
    with open_pdf(src) as doc:
        sel = _pages(spec, doc.page_count)
        if not sel:
            raise ValueError("No pages selected")
        doc.select(sel)
        _save(doc, dst)
        return doc.page_count


def thumbnails(src: str, width: int = 150, password: str = "") -> list[bytes]:
    """PNG thumbnails of every page (for the page grid)."""
    out = []
    with open_pdf(src, password) as doc:
        for pg in doc:
            z = width / max(pg.rect.width, 1)
            out.append(pg.get_pixmap(matrix=pymupdf.Matrix(z, z), alpha=False).tobytes("png"))
    return out


# ---------------------------------------------------------------- crop

def content_bbox(page: pymupdf.Page) -> pymupdf.Rect | None:
    box = pymupdf.Rect()
    for item in page.get_bboxlog():
        r = pymupdf.Rect(item[1])
        if r.is_empty or r.is_infinite or r.width > page.rect.width * 0.999 and r.height > page.rect.height * 0.999:
            continue
        box |= r
    return None if box.is_empty else box & page.rect


def crop_pdf(src: str, dst: str, margins: tuple[float, float, float, float] = (0, 0, 0, 0),
             pages: str = "", auto: bool = False, pad: float = 8, unit: str = "pt") -> int:
    """Crop pages. margins = (left, top, right, bottom) removed; unit pt|mm|percent. auto trims white margins."""
    with open_pdf(src) as doc:
        for i in _pages(pages, doc.page_count):
            pg = doc[i]
            r = pg.rect
            if auto:
                bb = content_bbox(pg)
                if not bb:
                    continue
                new = pymupdf.Rect(bb.x0 - pad, bb.y0 - pad, bb.x1 + pad, bb.y1 + pad) & r
            else:
                l, t, rt, b = margins
                if unit == "mm":
                    l, t, rt, b = (v * 72 / 25.4 for v in (l, t, rt, b))
                elif unit == "percent":
                    l, rt = l * r.width / 100, rt * r.width / 100
                    t, b = t * r.height / 100, b * r.height / 100
                new = pymupdf.Rect(r.x0 + l, r.y0 + t, r.x1 - rt, r.y1 - b)
            if new.width < 10 or new.height < 10:
                raise ValueError("The crop leaves nothing of the page")
            pg.set_cropbox(new)
        _save(doc, dst)
        return doc.page_count


# ---------------------------------------------------------------- blank pages

def find_blank_pages(src: str, tolerance: float = 0.002, ignore_text: bool = False) -> list[int]:
    """Pages whose rendering is (nearly) white. tolerance = fraction of non-white pixels allowed."""
    blanks = []
    with open_pdf(src) as doc:
        for i, pg in enumerate(doc):
            if not ignore_text and pg.get_text().strip():
                continue
            pix = pg.get_pixmap(matrix=pymupdf.Matrix(0.35, 0.35), colorspace=pymupdf.csGRAY, alpha=False)
            data = pix.samples
            dark = sum(1 for v in data if v < 235)
            if dark / max(len(data), 1) <= tolerance:
                blanks.append(i)
    return blanks


def delete_blank_pages(src: str, dst: str, tolerance: float = 0.002) -> list[int]:
    blanks = find_blank_pages(src, tolerance)
    with open_pdf(src) as doc:
        if len(blanks) >= doc.page_count:
            raise ValueError("Every page looks blank")
        if blanks:
            doc.delete_pages(blanks)
        _save(doc, dst)
    return blanks


# ---------------------------------------------------------------- header / footer / Bates / watermark / background

def header_footer(src: str, dst: str, left: str = "", center: str = "", right: str = "", where: str = "bottom",
                  size: float = 10, color="#444444", margin: float = 28, pages: str = "",
                  first_number: int = 1, bates: tuple[str, int, int] | None = None) -> int:
    """Stamp text on every page. Tokens: {page} {pages} {file} {date} {bates}."""
    name = Path(src).stem
    with open_pdf(src) as doc:
        n = doc.page_count
        for i in _pages(pages, n):
            pg = doc[i]
            r = pg.rect
            h = size * 2.2
            y0 = r.y0 + margin - size * 0.6 if where == "top" else r.y1 - margin - h + size * 0.6
            w = (r.width - 2 * margin) / 3
            for k, (txt, al) in enumerate(((left, "left"), (center, "center"), (right, "right"))):
                if not txt.strip():
                    continue
                t = _tokens(txt, i + first_number - 1, n + first_number - 1, name, bates)
                box = pymupdf.Rect(r.x0 + margin + k * w, y0, r.x0 + margin + (k + 1) * w, y0 + h)
                put_text(pg, box, t, size, color, align=al)
        _save(doc, dst)
        return n


def _text_image(text: str, size: float, color, bold: bool, pad: int = 12) -> "Image.Image":
    """Render text (shaped, RTL) to a transparent RGBA image using MuPDF."""
    from PIL import Image
    w = max(200, int(size * (len(text) * 0.75 + 4)))
    h = int(size * 2.6)
    with pymupdf.open() as tmp:
        pg = tmp.new_page(width=w, height=h)
        put_text(pg, pymupdf.Rect(0, 0, w, h), text, size, color, bold=bold, align="center")
        pix = pg.get_pixmap(matrix=pymupdf.Matrix(3, 3), alpha=True)
        img = Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGBA")
    box = img.getbbox()
    if box:
        img = img.crop((max(box[0] - pad, 0), max(box[1] - pad, 0), box[2] + pad, box[3] + pad))
    return img


def _fade(img: "Image.Image", opacity: float) -> "Image.Image":
    img = img.convert("RGBA")
    a = img.getchannel("A").point(lambda v: int(v * max(0.0, min(opacity, 1.0))))
    img.putalpha(a)
    return img


def watermark(src: str, dst: str, text: str = "", image: str = "", size: float = 64, color="#999999",
              opacity: float = 0.3, angle: float = 45, pages: str = "", tile: bool = False,
              scale: float = 0.6, behind: bool = False) -> int:
    """Text or image watermark, rotated by any angle, optionally tiled across the page."""
    from PIL import Image
    if not text.strip() and not image:
        raise ValueError("Enter watermark text or choose an image")
    base = Image.open(image).convert("RGBA") if image else _text_image(text, size, color, True)
    base = _fade(base.rotate(angle, expand=True, resample=Image.BICUBIC), opacity)
    buf = io.BytesIO()
    base.save(buf, "PNG")
    png = buf.getvalue()
    with open_pdf(src) as doc:
        for i in _pages(pages, doc.page_count):
            pg = doc[i]
            r = pg.rect
            pw = r.width * scale if not tile else r.width * 0.34
            ph = pw * base.height / base.width
            if ph > r.height * 0.9:
                ph = r.height * 0.9
                pw = ph * base.width / base.height
            if tile:
                y = r.y0 + 20
                row = 0
                while y < r.y1:
                    x = r.x0 + 20 + (pw / 2 if row % 2 else 0)
                    while x < r.x1:
                        pg.insert_image(pymupdf.Rect(x, y, x + pw, y + ph), stream=png, overlay=not behind)
                        x += pw * 1.15
                    y += ph * 1.15
                    row += 1
            else:
                c = r.tl + (r.br - r.tl) / 2
                pg.insert_image(pymupdf.Rect(c.x - pw / 2, c.y - ph / 2, c.x + pw / 2, c.y + ph / 2),
                                stream=png, overlay=not behind)
        _save(doc, dst)
        return doc.page_count


def background(src: str, dst: str, color="", image: str = "", pages: str = "") -> int:
    """Paint a colour or picture behind the page content."""
    if not color and not image:
        raise ValueError("Choose a colour or a picture")
    with open_pdf(src) as doc:
        for i in _pages(pages, doc.page_count):
            pg = doc[i]
            if image:
                pg.insert_image(pg.rect, filename=image, overlay=False, keep_proportion=False)
            else:
                pg.draw_rect(pg.rect, color=None, fill=_rgb(color), overlay=False)
        _save(doc, dst)
        return doc.page_count


# ---------------------------------------------------------------- compare

def _words(text: str) -> list[str]:
    return core.normalize_persian(text).split()


def compare_text(a: str, b: str) -> list[dict]:
    """Per-page text comparison. Returns [{page, same, ratio, html}] (html marks removed/added words)."""
    rows = []
    with open_pdf(a) as da, open_pdf(b) as db:
        for i in range(max(da.page_count, db.page_count)):
            wa = _words(da[i].get_text()) if i < da.page_count else []
            wb = _words(db[i].get_text()) if i < db.page_count else []
            sm = difflib.SequenceMatcher(None, wa, wb, autojunk=False)
            parts = []
            for op, i1, i2, j1, j2 in sm.get_opcodes():
                if op == "equal":
                    parts.append(_html.escape(" ".join(wa[i1:i2])))
                else:
                    if i2 > i1:
                        parts.append('<span style="background:#ffc9c9;text-decoration:line-through">'
                                     + _html.escape(" ".join(wa[i1:i2])) + "</span>")
                    if j2 > j1:
                        parts.append('<span style="background:#b2f2bb">' + _html.escape(" ".join(wb[j1:j2])) + "</span>")
            rows.append({"page": i, "same": wa == wb, "ratio": sm.ratio() if (wa or wb) else 1.0,
                         "html": " ".join(parts)})
    return rows


def compare_visual(a: str, b: str, index_a: int, index_b: int | None = None, zoom: float = 1.2) -> tuple[bytes, float]:
    """PNG with the differences between two pages in red, and the share of changed pixels."""
    from PIL import Image, ImageChops
    index_b = index_a if index_b is None else index_b
    with open_pdf(a) as da, open_pdf(b) as db:
        def render(doc, idx, size=None):
            if idx >= doc.page_count:
                return Image.new("RGB", size or (600, 800), "white")
            pix = doc[idx].get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
            return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        ia = render(da, index_a)
        ib = render(db, index_b, ia.size).resize(ia.size)
    diff = ImageChops.difference(ia.convert("L"), ib.convert("L")).point(lambda v: 255 if v > 40 else 0)
    changed = sum(1 for v in diff.getdata() if v) / max(diff.width * diff.height, 1)
    out = Image.blend(ia, Image.new("RGB", ia.size, "white"), 0.6)
    out.paste(Image.new("RGB", ia.size, (230, 30, 30)), mask=diff.resize(ia.size))
    buf = io.BytesIO()
    out.save(buf, "PNG")
    return buf.getvalue(), changed


def compare_report(a: str, b: str, dst: str) -> tuple[int, int]:
    """Write a PDF report listing every page that differs. Returns (pages_compared, pages_different)."""
    rows = compare_text(a, b)
    diff = [r for r in rows if not r["same"]]
    rtl = any(core.is_rtl_text(re.sub(r"<[^>]+>", "", r["html"])) for r in rows)
    body = [f"<h1>{_html.escape(Path(a).name)} ↔ {_html.escape(Path(b).name)}</h1>",
            f"<p>{len(diff)} of {len(rows)} pages differ. "
            '<span style="background:#ffc9c9">removed</span> · <span style="background:#b2f2bb">added</span></p>']
    for r in diff:
        body.append(f"<h2>Page {r['page'] + 1} — {r['ratio'] * 100:.0f}% similar</h2><p>{r['html']}</p>")
    if not diff:
        body.append("<p>No text differences found.</p>")
    core.html_to_pdf("<html><body>" + "".join(body) + "</body></html>", dst, "a4", 10, rtl=rtl)
    return len(rows), len(diff)


# ---------------------------------------------------------------- export

def to_images(src: str, out_dir: str, fmt: str = "png", dpi: int = 150, pages: str = "",
              progress: Progress = None) -> list[str]:
    """Save pages as png / jpg / tiff / bmp / webp pictures."""
    from PIL import Image
    fmt = fmt.lower().replace("jpeg", "jpg")
    os.makedirs(out_dir, exist_ok=True)
    files = []
    stem = Path(src).stem
    with open_pdf(src) as doc:
        sel = _pages(pages, doc.page_count)
        for k, i in enumerate(sel):
            pix = doc[i].get_pixmap(dpi=dpi, alpha=False)
            path = core.unique_path(Path(out_dir) / f"{stem}_{i + 1:03d}.{fmt}")
            if fmt == "png":
                pix.save(str(path))
            else:
                img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                img.save(path, {"jpg": "JPEG", "tiff": "TIFF"}.get(fmt, fmt.upper()), dpi=(dpi, dpi),
                         **({"quality": 90} if fmt == "jpg" else {}))
            files.append(str(path))
            if progress:
                progress(k + 1, len(sel), path.name)
    return files


def _blocks(page: pymupdf.Page):
    """Reading-order blocks: ('text', text, size, bold) or ('image', bytes)."""
    d = page.get_text("dict", sort=True)
    for b in d["blocks"]:
        if b["type"] == 0:
            lines = []
            sizes, bold = [], False
            for ln in b["lines"]:
                t = "".join(s["text"] for s in ln["spans"])
                if t.strip():
                    lines.append(t)
                for s in ln["spans"]:
                    if s["text"].strip():
                        sizes.append(s["size"])
                        bold = bold or bool(s["flags"] & 16)
            if lines:
                yield "text", lines, (sum(sizes) / len(sizes) if sizes else 11), bold
        elif b["type"] == 1 and b.get("image"):
            yield "image", b["image"], b.get("ext", "png"), b["bbox"]


def _join_lines(lines: list[str]) -> str:
    out = ""
    for ln in lines:
        ln = ln.rstrip()
        if out and out.endswith("-") and not out.endswith(" -"):
            out = out[:-1] + ln.lstrip()
        else:
            out = (out + " " + ln.lstrip()) if out else ln
    return core.normalize_persian(out.strip())


def _docx_layout(src: str, dst: str, pages: str) -> int:
    """PDF -> Word with columns, tables and positions kept (pdf2docx). Left-to-right documents."""
    from pdf2docx import Converter
    with open_pdf(src) as d:
        sel = _pages(pages, d.page_count)
    cv = Converter(src)
    try:
        cv.convert(dst, pages=sel)
    finally:
        cv.close()
    return len(sel)


def to_docx(src: str, dst: str, pages: str = "", scanned_images: bool = True,
            progress: Progress = None, engine: str = "auto") -> int:
    """PDF -> Word.

    engine: "layout" keeps columns/tables (pdf2docx; no right-to-left support), "flow" writes simple paragraphs
    (RTL aware, Persian safe), "auto" uses layout for left-to-right documents and flow for Persian/Arabic.
    """
    if engine != "flow":
        use_layout = engine == "layout"
        if engine == "auto":
            with open_pdf(src) as d:
                sel = _pages(pages, d.page_count)[:6]
                sample = " ".join(d[i].get_text() for i in sel)
            use_layout = not core.is_rtl_text(sample)
        if use_layout:
            try:
                return _docx_layout(src, dst, pages)
            except Exception:
                if engine == "layout":
                    raise
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Inches

    out = docx.Document()
    sec = out.sections[0]
    with open_pdf(src) as doc:
        sel = _pages(pages, doc.page_count)
        r0 = doc[sel[0]].rect
        sec.page_width, sec.page_height = Pt(r0.width), Pt(r0.height)
        for s in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, s, Pt(54))
        usable = (r0.width - 108) / 72
        for k, i in enumerate(sel):
            pg = doc[i]
            if k:
                out.add_page_break()
            blocks = list(_blocks(pg))
            has_text = any(b[0] == "text" for b in blocks)
            if not has_text and scanned_images:
                pix = pg.get_pixmap(dpi=150, alpha=False)
                out.add_picture(io.BytesIO(pix.tobytes("png")), width=Inches(usable))
                continue
            for b in blocks:
                if b[0] == "image":
                    try:
                        w = min(max((b[3][2] - b[3][0]) / 72, 0.3), usable)
                        out.add_picture(io.BytesIO(b[1]), width=Inches(w))
                    except Exception:
                        pass
                    continue
                _, lines, size, bold = b
                text = _join_lines(lines)
                rtl = core.is_rtl_text(text)
                p = out.add_paragraph()
                run = p.add_run(text)
                run.font.size = Pt(round(min(max(size, 6), 40) * 2) / 2)
                run.bold = bold or size > 15
                if rtl:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    ppr = p._p.get_or_add_pPr()
                    ppr.append(OxmlElement("w:bidi"))
                    rpr = run._r.get_or_add_rPr()
                    rpr.append(OxmlElement("w:rtl"))
                    fonts = rpr.find(qn("w:rFonts"))
                    if fonts is None:
                        fonts = OxmlElement("w:rFonts")
                        rpr.insert(0, fonts)
                    fonts.set(qn("w:cs"), "Tahoma")
                    szcs = OxmlElement("w:szCs")
                    szcs.set(qn("w:val"), str(int(round(min(max(size, 6), 40) * 2))))
                    rpr.append(szcs)
            if progress:
                progress(k + 1, len(sel), f"page {i + 1}")
        count = len(sel)
    out.save(dst)
    return count


def to_xlsx(src: str, dst: str, pages: str = "") -> int:
    """Detect tables and write each to its own sheet. Returns the number of tables."""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    found = 0
    with open_pdf(src) as doc:
        for i in _pages(pages, doc.page_count):
            tables = []
            for strat in ("lines", "text"):  # ruled tables first, then whitespace-aligned ones
                try:
                    tables = doc[i].find_tables(strategy=strat).tables
                except Exception:
                    tables = []
                if tables:
                    break
            for t_i, tab in enumerate(tables):
                rows = tab.extract()
                if not rows:
                    continue
                found += 1
                ws = wb.create_sheet(f"p{i + 1}_t{t_i + 1}"[:31])
                if any(core.is_rtl_text(str(c or "")) for r in rows for c in r):
                    ws.sheet_view.rightToLeft = True
                for r in rows:
                    ws.append([core.normalize_persian(str(c)).strip() if c is not None else "" for c in r])
    if not found:
        raise ValueError("No tables were found in this PDF")
    wb.save(dst)
    return found


def to_html_flow(src: str, pages: str = "", images: bool = True) -> str:
    """Simple flowing HTML (headings guessed from font size), pictures embedded as data URIs."""
    parts = []
    with open_pdf(src) as doc:
        sizes = [b[2] for p in doc for b in _blocks(p) if b[0] == "text"]
        body = sorted(sizes)[len(sizes) // 2] if sizes else 11
        rtl_all = core.is_rtl_text(" ".join(doc[i].get_text() for i in range(min(doc.page_count, 5))))
        for i in _pages(pages, doc.page_count):
            parts.append(f'<section id="p{i + 1}">')
            for b in _blocks(doc[i]):
                if b[0] == "image":
                    if images:
                        parts.append(f'<p><img src="data:image/{b[2]};base64,{base64.b64encode(b[1]).decode()}"></p>')
                    continue
                text = _html.escape(_join_lines(b[1]))
                d = ' dir="rtl"' if core.is_rtl_text(text) else ""
                tag = "h1" if b[2] > body * 1.7 else "h2" if b[2] > body * 1.35 else "h3" if (b[2] > body * 1.12 and b[3]) else "p"
                parts.append(f"<{tag}{d}>{text}</{tag}>")
            parts.append("</section>")
    d = ' dir="rtl"' if rtl_all else ""
    return (f'<!DOCTYPE html><html{d}><head><meta charset="utf-8"><title>{_html.escape(Path(src).stem)}</title>'
            f'<style>body{{font-family:Vazirmatn,Tahoma,Segoe UI,sans-serif;max-width:46em;margin:2em auto;line-height:1.7}}'
            f'img{{max-width:100%}}section{{border-bottom:1px solid #ccc;padding-bottom:1em}}</style></head>'
            f'<body>{"".join(parts)}</body></html>')


def export_text_format(src: str, dst: str, fmt: str, pages: str = "") -> str:
    """fmt: html | rtf | epub | txt. rtf/epub need Pandoc (bundled)."""
    fmt = fmt.lower()
    if fmt == "txt":
        core.extract_text(src, dst, "txt")
        return dst
    html = to_html_flow(src, pages, images=fmt != "rtf")
    if fmt == "html":
        Path(dst).write_text(html, encoding="utf-8")
        return dst
    pandoc = core.find_pandoc()
    if not pandoc:
        raise RuntimeError("Pandoc was not found")
    import subprocess
    tmp = os.path.join(tempfile.gettempdir(), f"pdftk_exp_{os.getpid()}.html")
    Path(tmp).write_text(html, encoding="utf-8")
    try:
        args = [pandoc, tmp, "-f", "html", "-o", dst, "--metadata", f"title={Path(src).stem}"]
        if fmt == "epub":
            args += ["--metadata", "lang=" + ("fa" if core.is_rtl_text(html) else "en")]
        r = subprocess.run(args, capture_output=True, text=True, creationflags=core.NO_WINDOW)
        if r.returncode:
            raise RuntimeError(r.stderr.strip()[:300] or "Pandoc failed")
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass
    return dst


# ---------------------------------------------------------------- security

PERMISSIONS = {"print": pymupdf.PDF_PERM_PRINT | pymupdf.PDF_PERM_PRINT_HQ, "copy": pymupdf.PDF_PERM_COPY,
               "edit": pymupdf.PDF_PERM_MODIFY | pymupdf.PDF_PERM_ASSEMBLE,
               "annotate": pymupdf.PDF_PERM_ANNOTATE | pymupdf.PDF_PERM_FORM}


def protect(src: str, dst: str, user_pw: str = "", owner_pw: str = "", allow: dict | None = None,
            password: str = "") -> None:
    """Encrypt with AES-256. user_pw opens the file; owner_pw lifts the restrictions in `allow`."""
    if not user_pw and not owner_pw:
        raise ValueError("Enter a password")
    allow = allow or {"print": True, "copy": True, "edit": True, "annotate": True}
    perm = pymupdf.PDF_PERM_ACCESSIBILITY
    for k, flag in PERMISSIONS.items():
        if allow.get(k, True):
            perm |= flag
    with open_pdf(src, password) as doc:
        _save(doc, dst, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw=user_pw,
              owner_pw=owner_pw or user_pw, permissions=perm)


def unlock(src: str, dst: str, password: str) -> None:
    with open_pdf(src, password) as doc:
        _save(doc, dst, encryption=pymupdf.PDF_ENCRYPT_NONE)


def is_encrypted(src: str) -> bool:
    with pymupdf.open(src) as doc:
        return bool(doc.needs_pass or doc.is_encrypted)


# ---------------------------------------------------------------- hidden data and redaction

HIDDEN_OPTIONS = ("metadata", "comments", "attachments", "javascript", "links", "bookmarks", "form_data",
                  "hidden_text", "thumbnails")


def remove_hidden(src: str, dst: str, options: dict | None = None) -> list[str]:
    """Scrub the chosen kinds of hidden information. Returns what was removed.

    'hidden_text' also deletes invisible OCR text layers, so it is off by default.
    """
    o = {k: True for k in HIDDEN_OPTIONS}
    o["hidden_text"] = False
    o.update(options or {})
    removed = []
    with open_pdf(src) as doc:
        if o["comments"]:
            for pg in doc:
                for a in list(pg.annots() or []):
                    if a.type[0] != pymupdf.PDF_ANNOT_WIDGET:
                        pg.delete_annot(a)
            removed.append("comments")
        if o["bookmarks"]:
            doc.set_toc([])
            removed.append("bookmarks")
        doc.scrub(metadata=o["metadata"], xml_metadata=o["metadata"], attached_files=o["attachments"],
                  embedded_files=o["attachments"], javascript=o["javascript"], remove_links=o["links"],
                  hidden_text=o["hidden_text"], thumbnails=o["thumbnails"], reset_fields=o["form_data"],
                  reset_responses=o["comments"], clean_pages=True, redactions=True)
        removed += [k for k in ("metadata", "attachments", "javascript", "links", "hidden_text",
                                "thumbnails", "form_data") if o[k]]
        _save(doc, dst)
    return removed


PATTERNS = {
    "email": r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+",
    "phone": r"(?:\+?\d[\d\s().-]{7,}\d)",
    "url": r"https?://[^\s]+|www\.[^\s]+",
    "card": r"\b(?:\d[ -]?){13,19}\b",
    "iban": r"\b[A-Z]{2}\d{2}[A-Z0-9]{11,30}\b",
    "digits_fa": r"[۰-۹]{6,}",
}


def find_terms(src: str, pattern: str, password: str = "") -> list[str]:
    """Distinct strings matching a regular expression, for the redaction list."""
    rx = re.compile(PATTERNS.get(pattern, pattern), re.I)
    seen = []
    with open_pdf(src, password) as doc:
        for pg in doc:
            for m in rx.finditer(pg.get_text()):
                s = m.group(0).strip()
                if s and s not in seen:
                    seen.append(s)
    return seen


def redact(src: str, dst: str, terms: list[str] | None = None, areas: dict[int, list[tuple]] | None = None,
           fill="#000000", pages: str = "", whole_words: bool = False) -> int:
    """Permanently remove text (and underlying pixels/pictures) for each term and rectangle.

    areas = {page_index: [(x0, y0, x1, y1), ...]} in PDF points. Returns the number of marks applied.
    """
    terms = [t for t in (terms or []) if t.strip()]
    if not terms and not areas:
        raise ValueError("Nothing to redact")
    count = 0
    with open_pdf(src) as doc:
        sel = set(_pages(pages, doc.page_count))
        for i, pg in enumerate(doc):
            if i not in sel and not (areas and i in areas):
                continue
            if i in sel:
                for t in terms:
                    for r in find_rects(pg, t):
                        pg.add_redact_annot(r, fill=_rgb(fill))
                        count += 1
            for r in (areas or {}).get(i, []):
                pg.add_redact_annot(pymupdf.Rect(r), fill=_rgb(fill))
                count += 1
            pg.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_PIXELS)
        doc.scrub(metadata=False, xml_metadata=False, attached_files=False, embedded_files=False,
                  javascript=False, remove_links=False, hidden_text=False, thumbnails=False,
                  reset_fields=False, reset_responses=False, clean_pages=True, redactions=True)
        _save(doc, dst)
    return count


# ---------------------------------------------------------------- editing

def add_text(src: str, dst: str, page: int, rect: tuple, text: str, size: float = 12, color="#000000",
             bold: bool = False, align: str = "left") -> bool:
    with open_pdf(src) as doc:
        ok = put_text(doc[page], pymupdf.Rect(rect), text, size, color, bold, align)
        _save(doc, dst)
    return ok


def add_image(src: str, dst: str, page: int, rect: tuple, image: str) -> None:
    with open_pdf(src) as doc:
        doc[page].insert_image(pymupdf.Rect(rect), filename=image, keep_proportion=True)
        _save(doc, dst)


def list_images(src: str, page: int) -> list[dict]:
    """Pictures on a page: [{xref, rect, width, height}]."""
    out = []
    with open_pdf(src) as doc:
        pg = doc[page]
        for info in pg.get_images(full=True):
            xref = info[0]
            for r in pg.get_image_rects(xref):
                out.append({"xref": xref, "rect": tuple(r), "width": info[2], "height": info[3]})
    return out


def replace_image(src: str, dst: str, page: int, xref: int, new_image: str) -> None:
    with open_pdf(src) as doc:
        doc[page].replace_image(xref, filename=new_image)
        _save(doc, dst)


def delete_image(src: str, dst: str, page: int, xref: int) -> None:
    with open_pdf(src) as doc:
        doc[page].delete_image(xref)
        _save(doc, dst)


def replace_text(src: str, dst: str, old: str, new: str, pages: str = "") -> int:
    """Find `old` and put `new` in its place with similar size and colour. Returns replacements made."""
    if not old.strip():
        raise ValueError("Enter the text to find")
    n = 0
    with open_pdf(src) as doc:
        for i in _pages(pages, doc.page_count):
            pg = doc[i]
            jobs = []
            for r in find_rects(pg, old):
                size, color, bold = r.height * 0.78, (0, 0, 0), False
                for b in pg.get_text("dict", clip=r)["blocks"]:
                    for ln in b.get("lines", []):
                        for s in ln["spans"]:
                            if s["text"].strip():
                                size, bold = s["size"], bool(s["flags"] & 16)
                                c = s["color"]
                                color = ((c >> 16 & 255) / 255, (c >> 8 & 255) / 255, (c & 255) / 255)
                                break
                jobs.append((r, size, color, bold))
            for r, *_ in jobs:
                pg.add_redact_annot(r, fill=(1, 1, 1))
            if jobs:
                pg.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE)
            for r, size, color, bold in jobs:
                need = len(new) * size * 0.62
                if core.is_rtl_text(new):  # grow to the left, keep the right edge
                    wide = pymupdf.Rect(min(r.x0, r.x1 - need) - 2, r.y0 - 1, r.x1 + 1, r.y1 + 2)
                else:
                    wide = pymupdf.Rect(r.x0 - 1, r.y0 - 1, max(r.x1, r.x0 + need) + 2, r.y1 + 2)
                put_text(pg, wide, new, size, color, bold)
                n += 1
        _save(doc, dst)
    return n


# ---------------------------------------------------------------- scan cleanup (OpenCV) and repair (pikepdf)

def estimate_skew(gray) -> float:
    """Skew angle of a scanned page in degrees (positive = text runs downhill), found by projection profiles."""
    import cv2
    import numpy as np
    h, w = gray.shape[:2]
    scale = 900.0 / max(h, w)
    small = cv2.resize(gray, (max(int(w * scale), 1), max(int(h * scale), 1)), interpolation=cv2.INTER_AREA)
    _t, bw = cv2.threshold(small, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)
    if cv2.countNonZero(bw) < 50:
        return 0.0
    cy, cx = bw.shape[0] / 2, bw.shape[1] / 2

    def score(angle: float) -> float:
        m = cv2.getRotationMatrix2D((cx, cy), angle, 1.0)
        r = cv2.warpAffine(bw, m, (bw.shape[1], bw.shape[0]), flags=cv2.INTER_NEAREST)
        prof = r.sum(axis=1).astype(np.float64)
        return float(np.var(np.diff(prof)))

    best = max(np.arange(-10, 10.01, 0.5), key=score)
    best = max(np.arange(best - 0.5, best + 0.51, 0.1), key=score)
    return float(best)


def _clean_image(img, mode: str, deskew: bool, denoise: str, shadows: bool):
    import cv2
    import numpy as np
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY) if img.ndim == 3 else img
    angle = estimate_skew(gray) if deskew else 0.0
    colour = mode == "color" and img.ndim == 3
    if abs(angle) >= 0.15:
        h, w = gray.shape[:2]
        m = cv2.getRotationMatrix2D((w / 2, h / 2), angle, 1.0)
        white = (255, 255, 255)
        img = cv2.warpAffine(img, m, (w, h), flags=cv2.INTER_CUBIC, borderValue=white if img.ndim == 3 else 255)
        gray = cv2.warpAffine(gray, m, (w, h), flags=cv2.INTER_CUBIC, borderValue=255)
    if colour:                                  # keep the colours: only straighten, flatten the light and denoise
        out = img
        if shadows:
            bg = cv2.medianBlur(cv2.resize(gray, None, fx=0.25, fy=0.25, interpolation=cv2.INTER_AREA), 21)
            bg = cv2.resize(bg, (gray.shape[1], gray.shape[0]), interpolation=cv2.INTER_LINEAR)
            gain = (255.0 / np.maximum(bg, 1).astype(np.float32))[:, :, None]
            out = np.clip(out.astype(np.float32) * gain, 0, 255).astype(np.uint8)
        if denoise == "strong":
            out = cv2.fastNlMeansDenoisingColored(out, None, 8, 8, 7, 21)
        elif denoise == "light":
            out = cv2.medianBlur(out, 3)
        return out, angle
    if shadows:
        k = max(int(min(gray.shape) / 30) | 1, 15)
        bg = cv2.medianBlur(cv2.resize(gray, None, fx=0.25, fy=0.25, interpolation=cv2.INTER_AREA), max(k // 4 | 1, 5))
        bg = cv2.resize(bg, (gray.shape[1], gray.shape[0]), interpolation=cv2.INTER_LINEAR)
        gray = cv2.divide(gray, np.maximum(bg, 1), scale=255)
    if denoise == "light":
        gray = cv2.medianBlur(gray, 3)
    elif denoise == "strong":
        gray = cv2.fastNlMeansDenoising(gray, None, 12, 7, 21)
    if mode == "bw":
        out = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C, cv2.THRESH_BINARY, 41, 15)
        return out, angle
    return gray, angle


def clean_scan(src: str, dst: str, mode: str = "gray", deskew: bool = True, denoise: str = "light",
               shadows: bool = True, dpi: int = 300, pages: str = "", progress: Progress = None) -> list[float]:
    """Straighten, de-noise and clean scanned pages. Output pages are pictures: run OCR afterwards.

    mode: gray | bw (black & white, smallest files) | color. Returns the skew angle found on each page.
    """
    import cv2
    import numpy as np
    angles = []
    with open_pdf(src) as doc, pymupdf.open() as out:
        sel = _pages(pages, doc.page_count)
        for k, i in enumerate(sel):
            pg = doc[i]
            pix = pg.get_pixmap(dpi=dpi, colorspace=pymupdf.csRGB if mode == "color" else pymupdf.csGRAY, alpha=False)
            arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
            img = cv2.cvtColor(arr, cv2.COLOR_RGB2BGR) if pix.n == 3 else arr[:, :, 0].copy()
            res, angle = _clean_image(img, mode, deskew, denoise, shadows)
            angles.append(angle)
            ok, buf = cv2.imencode(".png" if mode == "bw" else ".jpg", res,
                                   [] if mode == "bw" else [cv2.IMWRITE_JPEG_QUALITY, 88])
            page = out.new_page(width=pg.rect.width, height=pg.rect.height)
            page.insert_image(page.rect, stream=buf.tobytes())
            if progress:
                progress(k + 1, len(sel), f"page {i + 1}")
        _save(out, dst)
    return angles


def _problems(pdf) -> list[str]:
    fn = getattr(pdf, "check_pdf_syntax", None) or getattr(pdf, "check", None)
    try:
        return [str(m) for m in (fn() if fn else [])]
    except Exception as e:
        return [str(e)]


def check_pdf(src: str) -> list[str]:
    """Structural problems found by qpdf (empty list = healthy)."""
    import pikepdf
    try:
        with pikepdf.open(src, attempt_recovery=True) as pdf:
            return _problems(pdf)
    except Exception as e:
        return [str(e)]


def repair_pdf(src: str, dst: str, linearize: bool = False, password: str = "") -> tuple[int, int, list[str]]:
    """Rebuild a damaged PDF and store it compactly. Returns (old_size, new_size, problems_found)."""
    import pikepdf
    problems = []
    with pikepdf.open(src, attempt_recovery=True, password=password) as pdf:
        problems = _problems(pdf)
        pdf.save(dst, linearize=linearize, object_stream_mode=pikepdf.ObjectStreamMode.generate,
                 compress_streams=True, recompress_flate=True)
    return os.path.getsize(src), os.path.getsize(dst), problems


# ---------------------------------------------------------------- batch

def _out_name(src: str, out_dir: str, suffix: str, ext: str = ".pdf") -> str:
    return str(core.unique_path(Path(out_dir) / f"{Path(src).stem}{suffix}{ext}"))


def _op_pdf(fn):
    return lambda src, out_dir, p: ([fn(src, _out_name(src, out_dir, p.get("suffix", "")), p)])


def _tuple(v, default):
    return tuple(v) if v else default


BATCH_OPS: dict[str, tuple[str, Callable]] = {
    # name: (default suffix, function(src, out_dir, params) -> list[str] of output files)
    "compress": ("_small", lambda s, d, p: [_do(core.compress_pdf, s, _out_name(s, d, p["suffix"]), p.get("level", "medium"), p.get("engine", "builtin"))]),
    "ocr": ("_ocr", lambda s, d, p: [_do(core.ocr_pdf, s, _out_name(s, d, p["suffix"]), p.get("lang", "eng"), dpi=p.get("dpi", 300))]),
    "protect": ("_locked", lambda s, d, p: [_do(protect, s, _out_name(s, d, p["suffix"]), p.get("user_pw", ""), p.get("owner_pw", ""), p.get("allow"))]),
    "unlock": ("_unlocked", lambda s, d, p: [_do(unlock, s, _out_name(s, d, p["suffix"]), p.get("password", ""))]),
    "crop": ("_cropped", lambda s, d, p: [_do(crop_pdf, s, _out_name(s, d, p["suffix"]), _tuple(p.get("margins"), (0, 0, 0, 0)), "", p.get("auto", False), 8, p.get("unit", "pt"))]),
    "delete_blank": ("_noblank", lambda s, d, p: [_do(delete_blank_pages, s, _out_name(s, d, p["suffix"]), p.get("tolerance", 0.002))]),
    "watermark": ("_wm", lambda s, d, p: [_do(watermark, s, _out_name(s, d, p["suffix"]), p.get("text", ""), p.get("image", ""), p.get("size", 64), p.get("color", "#999999"), p.get("opacity", 0.3), p.get("angle", 45))]),
    "header_footer": ("_hf", lambda s, d, p: [_do(header_footer, s, _out_name(s, d, p["suffix"]), p.get("left", ""), p.get("center", ""), p.get("right", ""), p.get("where", "bottom"), p.get("size", 10), p.get("color", "#444444"), 28, "", 1, p.get("bates"))]),
    "background": ("_bg", lambda s, d, p: [_do(background, s, _out_name(s, d, p["suffix"]), p.get("color", ""), p.get("image", ""))]),
    "remove_hidden": ("_clean", lambda s, d, p: [_do(remove_hidden, s, _out_name(s, d, p["suffix"]), p.get("options"))]),
    "to_images": ("", lambda s, d, p: to_images(s, str(Path(d) / Path(s).stem), p.get("fmt", "png"), p.get("dpi", 150))),
    "to_docx": ("", lambda s, d, p: [_do(to_docx, s, _out_name(s, d, p["suffix"], ".docx"))]),
    "clean_scan": ("_clean", lambda s, d, p: [_do(clean_scan, s, _out_name(s, d, p["suffix"]), p.get("mode", "gray"))]),
    "repair": ("_fixed", lambda s, d, p: [_do(repair_pdf, s, _out_name(s, d, p["suffix"]))]),
    "to_xlsx": ("", lambda s, d, p: [_do(to_xlsx, s, _out_name(s, d, p["suffix"], ".xlsx"))]),
    "to_html": ("", lambda s, d, p: [export_text_format(s, _out_name(s, d, p["suffix"], ".html"), "html")]),
    "to_text": ("", lambda s, d, p: [_do(core.extract_text, s, _out_name(s, d, p["suffix"], ".txt"), "txt")]),
    "convert": ("", lambda s, d, p: [core.convert_to_pdf(s, d, p.get("page_size", "a5"))]),
}


def _do(fn, src, dst, *args, **kw):
    fn(src, dst, *args, **kw)
    return dst


def batch_run(files: list[str], out_dir: str, op: str, params: dict | None = None,
              progress: Progress = None, cancel: Callable[[], bool] | None = None) -> list[tuple[str, bool, str]]:
    """Apply one operation to many files. Returns [(file, ok, message_or_output)]."""
    if op not in BATCH_OPS:
        raise ValueError(f"Unknown batch operation {op}")
    suffix, fn = BATCH_OPS[op]
    p = dict(params or {})
    p.setdefault("suffix", suffix)
    os.makedirs(out_dir, exist_ok=True)
    results = []
    for k, f in enumerate(files):
        if cancel and cancel():
            break
        if progress:
            progress(k, len(files), Path(f).name)
        try:
            outs = fn(f, out_dir, p)
            results.append((f, True, "; ".join(str(o) for o in outs if o) if outs else ""))
        except Exception as e:  # keep going: one bad file must not stop the batch
            results.append((f, False, str(e) or e.__class__.__name__))
    if progress:
        progress(len(files), len(files), "")
    return results
